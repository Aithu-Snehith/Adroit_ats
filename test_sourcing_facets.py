"""Sourcing "Filter saved profiles" dropdowns: College (Indian Bachelor's) / Passout year / Location.

Each dropdown lists only values present in saved profiles, with how many each gives (the other
dropdowns applied, not its own); several ticks in one dropdown = either; across dropdowns = all;
college + year must be the SAME degree; the "All saved profiles" tab has no education rule; location is read as a US state (full name or 2-letter) or
India; unknown values are refused; the Excel follows the dropdowns; they combine with the LinkedIn
checkboxes. Nothing is paid - only saved profiles are read. Throwaway SQLite DB, no network.
People are labelled test fixtures.
Run: python test_sourcing_facets.py
"""
import io
import os
import sys
import tempfile

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_facets.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import app as app_module  # noqa: E402
import education_filters as ef  # noqa: E402
import linkedin_ingest as li  # noqa: E402
import models  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


for loc, want in (("Dallas, Texas, United States", "Texas"), ("Austin, TX", "Texas"), ("San Jose, California", "California"),
                  ("Hyderabad, Telangana, India", "India"), ("Dallas-Fort Worth Metroplex", ""), ("", ""), ("Washington, DC", "District of Columbia")):
    check(ef.state_of(loc) == want, f"state_of({loc!r}) = {ef.state_of(loc)!r}, want {want!r}")

uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
ADMIN = {"id": uid(models.create_user("Facet Admin", "facet-admin@example.invalid", "Facet-1", role="Admin")), "name": "Facet Admin", "role": "Admin"}


def item(slug, first, college, year, loc, otw=None):
    return {"linkedinUrl": f"https://www.linkedin.com/in/{slug}.example.invalid", "firstName": first, "lastName": "Facettest",
            "headline": "Engineer", "location": {"linkedinText": loc}, "currentPosition": [{"companyName": "Co"}],
            "openToWork": otw, "experience": [],
            "education": [{"schoolName": college, "degree": "B.Tech", "endDate": {"year": year}},
                          {"schoolName": "University of Texas at Dallas", "degree": "MS", "endDate": {"year": year + 3}}]}


li.ingest_profiles(li.ApifyProfileProvider().to_profiles([
    item("f-1", "Anu", "JNTUH", 2019, "Dallas, Texas", otw=True), item("f-2", "Bhanu", "JNTUH", 2022, "San Jose, California, United States"),
    item("f-3", "Chandu", "CBIT", 2022, "Austin, TX"), item("f-4", "Divya", "CBIT", 2019, "Dallas-Fort Worth Metroplex")]), "apify", ADMIN["id"])

client = app_module.app.test_client()
with client.session_transaction() as s:
    s["user"] = {**ADMIN, "email": "facet-admin@example.invalid"}


def run(**q):
    r = client.get("/api/education/search", query_string={"filter": "P", "passout_year": "All", **q})
    d = r.get_json()
    return r.status_code, sorted(x["name"].split()[0] for x in d.get("results", [])), d


code, names, d = run()
check(code == 200 and names == ["Anu", "Bhanu", "Chandu", "Divya"], f"all four saved: {code} {names} {d.get('error')}")
f = d["facets"]
cols = {o["label"]: (o["value"], o["count"]) for o in f["colleges"]}
jntuh = next((v for lbl, (v, n) in cols.items() if "jawaharlal" in lbl.lower() or "jntu" in lbl.lower()), None)
cbit = next((v for lbl, (v, n) in cols.items() if "chaitanya" in lbl.lower() or "cbit" in lbl.lower()), None)
check(jntuh and cbit and sorted(n for v, n in cols.values()) == [2, 2], f"college options with counts: {f['colleges']}")
check([(o["value"], o["count"]) for o in f["years"]] == [(2022, 2), (2019, 2)], f"year options, newest first: {f['years']}")
check({o["value"]: o["count"] for o in f["states"]} == {"Texas": 2, "California": 1} and f["location_not_listed"] == 1, f"location options: {f['states']} / {f['location_not_listed']}")
check(d["counts"]["matching"] == 4 and d["counts"]["total_scraped"] == 4, f"cards: {d['counts']}")

code, names, d = run(colleges=jntuh)
check(names == ["Anu", "Bhanu"], f"one college: {names}")
check({o["value"]: o["count"] for o in d["facets"]["years"]} == {2019: 1, 2022: 1}, f"years follow the chosen college: {d['facets']['years']}")
check(sorted(o["count"] for o in d["facets"]["colleges"]) == [2, 2], "college counts ignore their own selection (so you can add more)")
check(run(colleges=f"{jntuh},{cbit}")[1] == ["Anu", "Bhanu", "Chandu", "Divya"], "two colleges = either")
check(run(colleges=jntuh, years="2022")[1] == ["Bhanu"], "college + year")
check(run(colleges=f"{jntuh},{cbit}", years="2019")[1] == ["Anu", "Divya"], "two colleges + one year")
check(run(years="2022", states="Texas")[1] == ["Chandu"], "year + location (TX read as Texas)")
check(run(states="Texas,California")[1] == ["Anu", "Bhanu", "Chandu"], "two states = either")
check(run(states="Texas", open_to_work="1")[1] == ["Anu"], "dropdowns + LinkedIn checkboxes")
code, names, d = run(colleges=cbit, years="2019")
check(d["counts"]["matching"] == 1 and d["total"] == 1, f"Match card follows the dropdowns: {d['counts']}")
check(run(states="Atlantis")[0] == 400 and run(years="20x")[0] == 400, "unknown location / bad year refused")

x = client.get("/api/education/export-xlsx", query_string={"filter": "P", "passout_year": "All", "states": "Texas"})
ws = load_workbook(io.BytesIO(x.data)).active
check(x.status_code == 200 and ws.max_row - 1 == 2, f"Excel follows the dropdowns: {ws.max_row - 1} rows")

# ---- "All saved profiles": no education rule - a profile without a US Master's is still reachable
li.ingest_profiles(li.ApifyProfileProvider().to_profiles([{
    "linkedinUrl": "https://www.linkedin.com/in/f-5.example.invalid", "firstName": "Farah", "lastName": "Facettest",
    "headline": "QA Engineer", "location": {"linkedinText": "Houston, Texas"}, "currentPosition": [], "openToWork": True, "experience": [],
    "education": [{"schoolName": "JNTUH", "degree": "B.Tech", "endDate": {"year": 2012}}]}]), "apify", ADMIN["id"])
check(run()[1] == ["Anu", "Bhanu", "Chandu", "Divya"], "Passout tab unchanged (Farah has no US Master's)")
code, names, d = run(filter="ALL")
check(code == 200 and names == ["Anu", "Bhanu", "Chandu", "Divya", "Farah"], f"All saved profiles: {code} {names} {d.get('error')}")
check(d["counts"]["matching"] == d["counts"]["total_scraped"] == 5, f"All tab matches the Total card: {d['counts']}")
farah = next(x for x in d["results"] if x["name"].startswith("Farah"))
check("2012" in farah["indian_college"] and farah["us_masters"] == "", f"row shows her Bachelor's: {farah}")
check(run(filter="ALL", states="Texas")[1] == ["Anu", "Chandu", "Farah"], "All + Location")
check(run(filter="ALL", open_to_work="1")[1] == ["Anu", "Farah"], "All + Open to work")
check(2012 in [o["value"] for o in run(filter="ALL")[2]["facets"]["years"]], "year list includes years outside 2015-2023")
x = client.get("/api/education/export-xlsx", query_string={"filter": "ALL"})
check(x.status_code == 200 and load_workbook(io.BytesIO(x.data)).active.max_row - 1 == 5 and "All_saved_profiles" in x.headers.get("Content-Disposition", ""),
      "Excel of all saved profiles")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    sys.exit(1)
print("PASS: sourcing saved-profile dropdowns - college / year / location options with counts, either-within / all-across, same-degree college+year, state parsing, refusals, Excel, combine with LinkedIn checkboxes.")
