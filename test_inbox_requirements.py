"""Requirement emails in a consultant's Gmail -> tailored resume -> Gmail draft reply (inbox_requirements.py).

A vendor requirement (like "Virtusa is Hiring: DevOps SRE Engineers!") with a match >= 60% gets a draft
"Re: <their subject>" in their thread (to the Reply-To address), with the tailored resume saved as a
version and attached, and appears in the tracker; a 40% match gets no draft; LinkedIn / Naukri
notifications are skipped without AI; a personal email is "not a requirement"; mail from an open
application's recruiter is left to the Replies feature; a second scan adds nothing; an AI outage is
retried; INBOX read-only + BODY.PEEK; other recruiters can't see it; the window shows it. Gmail, the
AI and the Resume Optimizer are stubs - no network, no cost. Labelled test fixtures.
Run: python test_inbox_requirements.py
"""
import base64
import email
import email.policy
import io
import json
import os
import socket
import sys
import tempfile
import threading
from email.message import EmailMessage

tmp_dir = tempfile.mkdtemp()
os.environ["DB_PATH"] = os.path.join(tmp_dir, "test_reqs.db")
os.environ["DATABASE_URL"] = ""
os.environ["APIFY_API_TOKEN"] = ""
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import config  # noqa: E402

config.DB_PATH = os.environ["DB_PATH"]
config.DATABASE_URL = ""
import docx  # noqa: E402
import app as app_module  # noqa: E402
import inbox_requirements as iq  # noqa: E402
import models  # noqa: E402
import resume_bot  # noqa: E402

failures = []


def check(cond, msg):
    if not cond:
        failures.append(msg)


INBOX, appended, searches, fetches, selected = {}, [], [], [], []


def mail(n, frm, subject, body, msg_id, reply_to=""):
    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = frm, "visakh.reqtest@gmail.example", subject
    m["Date"] = "Mon, 12 Oct 2026 07:00:00 +0000"
    m["Message-ID"] = msg_id
    if reply_to:
        m["Reply-To"] = reply_to
    m.set_content(body)
    INBOX[str(n).encode()] = m.as_bytes()


class FakeIMAP:
    def __init__(self, host, port=None):
        pass

    def login(self, u, p):
        pass

    def select(self, box, readonly=False):
        selected.append(readonly)
        return "OK", [b"1"]

    def search(self, charset, *criteria):
        searches.append(criteria)
        if criteria and criteria[0] == "X-GM-RAW":
            return "OK", [b" ".join(sorted(INBOX, key=int))]
        return "OK", [b""]          # the Replies feature's FROM searches

    def fetch(self, mid, what):
        fetches.append(what)
        return "OK", [(mid + b" (BODY[] {1}", INBOX[mid]), b")"]

    def append(self, folder, flags, when, raw):
        appended.append(email.message_from_bytes(raw, policy=email.policy.default))

    def logout(self):
        pass


iq.imaplib.IMAP4_SSL = FakeIMAP
resume_bot.ai_configured = lambda: True
resume_bot.gemini_configured = lambda: False
ai_calls = []


def fake_generate(client, types, system, prompt, want_json, max_tokens):
    ai_calls.append(prompt[:200])
    if system == iq.EXTRACT_SYSTEM:
        low = prompt.lower()
        if "hiring" not in low and "requirement" not in low:
            return json.dumps({"is_requirement": False}), False, "stub"
        title = "Java Developer" if "java developer" in low else "DevOps SRE Engineer"
        return json.dumps({"is_requirement": True, "title": title, "client": "Virtusa", "location": "Hyderabad", "employment_type": "Full-time",
                           "skills": ["AWS", "Kubernetes"], "recruiter_name": "", "recruiter_email": "not-in-the-email@x.example",
                           "job_description": f"{title}: AWS, Kubernetes, Terraform, CI/CD. 5+ years."}), False, "stub"
    return json.dumps({"category": "other", "reason": ""}), False, "stub"


resume_bot._generate = fake_generate
d = docx.Document()
d.add_paragraph("VISAKH REQTEST - DevOps Engineer")
d.add_paragraph("AWS, Kubernetes, Terraform")
buf = io.BytesIO()
d.save(buf)
ORIGINAL = buf.getvalue()
d.add_paragraph("CI/CD with Jenkins")
buf2 = io.BytesIO()
d.save(buf2)
TAILORED = buf2.getvalue()
opt_calls = []


def fake_optimize(resume_text, jd_text, custom_instructions="", docx_bytes=None):
    opt_calls.append(jd_text[:40])
    if "java developer" in jd_text.lower():
        return {"ai_powered": True, "initial_match_percentage": 40, "optimized": False, "match_decision": "Reject (below 60%)"}
    return {"ai_powered": True, "initial_match_percentage": 78, "optimized": True, "docx_base64": base64.b64encode(TAILORED).decode(),
            "mandatory_matched_skills": ["AWS", "Kubernetes"], "match_decision": "Moderate enhancement"}


resume_bot.optimize_resume_for_jd = fake_optimize

uid = lambda u: u["id"] if isinstance(u, dict) else u  # noqa: E731
EMAIL, PW = "req-rec@example.invalid", "Req-1"
REC = {"id": uid(models.create_user("Ravi Reqtest", EMAIL, PW, role="Recruiter", region="India")), "name": "Ravi Reqtest", "role": "Recruiter"}
OTHER = {"id": uid(models.create_user("Other Reqtest", "req-other@example.invalid", "Req-2", role="Recruiter", region="India")), "name": "Other", "role": "Recruiter"}
cid = models.create_candidate("Visakh Reqtest", "visakh.reqtest@gmail.example", phone="+91 62351 79249", title="DevOps Engineer",
                              primary_skills="AWS, Kubernetes, Terraform", experience_years=10, assigned_user_id=REC["id"])
models.update_candidate(cid, country="India", gmail_account="visakh.reqtest@gmail.example", gmail_app_password="fake app pass not real")
models.save_resume_file(cid, "Visakh_Resume.docx", ORIGINAL)
# an existing application whose recruiter writes back -> belongs to the Replies feature
j0 = models.save_or_update_scraped_job({"title": "Old Job Reqtest", "company": "Acme", "source": "Dice", "recruiter_email": "known@acme.example",
                                        "url": "https://www.dice.com/job-detail/rq-0", "description": "x"})
models.mark_job_drafted(j0["id"] if isinstance(j0, dict) else j0, cid, "d-1", user_id=REC["id"], recruiter_email="known@acme.example")

mail(1, "Devops Engineer <rakesh@virtusa-vendor.example>", "Virtusa is Hiring: DevOps SRE Engineers!",
     "Hi, we are hiring DevOps SRE Engineers - AWS, Kubernetes. Please share your resume.", "<req1@virtusa>", reply_to="rakesh.r@virtusa-vendor.example")
mail(2, "Priya <priya@javavendor.example>", "Urgent requirement: Java Developer", "Java Developer requirement, Spring Boot, 8+ years.", "<req2@java>")
mail(3, "LinkedIn <messages-noreply@linkedin.com>", "The LinkedIn just messaged you", "1 new message awaits your response", "<n3@linkedin>")
mail(4, "NaukriGulf <info@naukrigulf.com>", "Welcome! The secret to a perfect job search is inside", "Welcome", "<n4@naukrigulf>")
mail(5, "Friend <friend@personal.example>", "Lunch tomorrow?", "Are you free for lunch tomorrow?", "<p5@friend>")
mail(6, "Known <known@acme.example>", "Re: Old Job", "Thanks, we will get back.", "<r6@acme>")

conn = models.get_db_connection()
s = iq.scan_candidate(conn, models.get_candidate_by_id(cid), REC["id"])
check(s["drafted"] == 1 and s["low_match"] == 1 and s["filtered"] == 2 and s["not_requirement"] == 1 and s["skipped_known"] == 1, f"outcomes: {s}")
check(all(selected) and fetches and all("PEEK" in f for f in fetches), "INBOX read-only + BODY.PEEK")
check(any(c and c[0] == "X-GM-RAW" and "category:primary" in c[1] for c in searches), "Primary tab only")
check(len([c for c in ai_calls if c.startswith("FROM:")]) == 3, f"no AI for filtered / known mail: {len(ai_calls)} calls")
check(len(appended) == 1, f"exactly one draft: {len(appended)}")
m = appended[0] if appended else None
check(m is not None and m["Subject"] == "Re: Virtusa is Hiring: DevOps SRE Engineers!" and m["In-Reply-To"] == "<req1@virtusa>"
      and m["To"] == "rakesh.r@virtusa-vendor.example", f"threaded draft to the Reply-To address: {m and (m['Subject'], m['In-Reply-To'], m['To'])}")
att = [p for p in m.iter_attachments()] if m else []
check(att and att[0].get_payload(decode=True) == TAILORED and att[0].get_filename().startswith("Visakh"), f"tailored resume attached: {[a.get_filename() for a in att]}")
body = m.get_body(preferencelist=("plain",)).get_content() if m else ""
check("C2C" not in body and "authorized to work" not in body, "India consultant: no US wording in the pitch")
check(body.startswith("Hi,") and "Hi Devops Engineer" not in body, f"sender name 'Devops Engineer' is a role, not a person: {body[:40]!r}")
check("Thank you for sharing the DevOps SRE Engineer requirement at Virtusa" in body and "came across your opening" not in body,
      f"answers their email: {body[:200]!r}")
check("AWS and Kubernetes" in body and "updated resume, tailored to this requirement" in body, f"skills from the match + tailored resume: {body!r}")
rows = {r["message_id"]: r for r in iq.list_for(conn, [cid])}
r1 = rows["<req1@virtusa>"]
check(r1["status"] == "drafted" and r1["match_pct"] == 78 and r1["resume_version_id"] and r1["job_title"] == "DevOps SRE Engineer", f"row: {r1}")
check(rows["<req2@java>"]["status"] == "low_match" and "40%" in rows["<req2@java>"]["reason"], "low match kept with reason")
check(rows["<n3@linkedin>"]["status"] == "filtered" and rows["<p5@friend>"]["status"] == "not_requirement" and "<r6@acme>" not in rows, "skips")
job = models.get_job_by_id(r1["job_id"])
check(job["source"] == "Email requirement" and job["recruiter_email"] == "rakesh.r@virtusa-vendor.example", f"job created: {job['source']} {job['recruiter_email']}")
cur = conn.cursor()
cur.execute("SELECT stage, resume_version_id, outreach_message_id FROM applications WHERE job_id = ? AND candidate_id = ?", (r1["job_id"], cid))
a = cur.fetchone()
check(a and a["stage"] == "Drafted" and a["resume_version_id"] == r1["resume_version_id"] and a["outreach_message_id"], "tracked as Draft created with the version")

# ---- second scan: nothing new
s2 = iq.scan_candidate(conn, models.get_candidate_by_id(cid), REC["id"])
check(s2["checked"] == 0 and len(appended) == 1 and len(opt_calls) == 2, f"second scan adds nothing: {s2} drafts={len(appended)}")

# ---- AI outage -> error, retried later
resume_bot.ai_configured = lambda: False
mail(7, "Sam <sam@vendor7.example>", "Hiring: Kubernetes Engineer", "We are hiring a Kubernetes engineer - AWS, Kubernetes.", "<req7@v7>")
iq.scan_candidate(conn, models.get_candidate_by_id(cid), REC["id"])
check({r["message_id"]: r for r in iq.list_for(conn, [cid])}["<req7@v7>"]["status"] == "error", "AI outage -> error")
resume_bot.ai_configured = lambda: True
iq.scan_candidate(conn, models.get_candidate_by_id(cid), REC["id"])
r7 = {r["message_id"]: r for r in iq.list_for(conn, [cid])}["<req7@v7>"]
check(r7["status"] == "drafted" and r7["attempts"] == 2 and len(appended) == 2, f"retried and drafted: {r7['status']} attempts={r7['attempts']}")
b7 = appended[-1].get_body(preferencelist=("plain",)).get_content() if len(appended) == 2 else ""
check(b7.startswith("Hi Sam,"), f"a real person's first name is used: {b7[:30]!r}")
j1_, j7_ = r1["job_id"], r7["job_id"]
check(j1_ != j7_ and models.get_job_by_id(j1_)["recruiter_email"] == "rakesh.r@virtusa-vendor.example"
      and models.get_job_by_id(j7_)["recruiter_email"] == "sam@vendor7.example",
      "same title + client from two vendors = two jobs, each keeps its own recruiter")
conn.close()

# ---- API + privacy
client = app_module.app.test_client()
with client.session_transaction() as s_:
    s_["user"] = {**REC, "email": EMAIL}
lst = client.get(f"/api/requirements?candidate_id={cid}").get_json()
check(len(lst["requirements"]) == 6 and lst["min_match"] == 60, f"listed: {len(lst['requirements'])}")
mail(8, "Ann <ann@vendor8.example>", "Requirement: SRE", "Hiring SRE - AWS, Kubernetes.", "<req8@v8>")
r = client.post("/api/requirements/scan", json={"candidate_id": cid}).get_json()
check(r["started"], f"scan started in the background: {r}")
for t in list(app_module._req_scans.values()):
    t.join(30)
check(any(x["message_id"] == "<req8@v8>" and x["status"] == "drafted" for x in client.get(f"/api/requirements?candidate_id={cid}").get_json()["requirements"]), "background scan drafted it")
with client.session_transaction() as s_:
    s_["user"] = {**OTHER, "email": "req-other@example.invalid"}
check(client.get(f"/api/requirements?candidate_id={cid}").status_code == 404 and client.get("/api/requirements").get_json()["requirements"] == [], "other recruiters see nothing")

# ---- the window (Edge)
from playwright.sync_api import sync_playwright  # noqa: E402
from werkzeug.serving import make_server  # noqa: E402

with socket.socket() as sk:
    sk.bind(("127.0.0.1", 0))
    port = sk.getsockname()[1]
server = make_server("127.0.0.1", port, app_module.app, threaded=True)
threading.Thread(target=server.serve_forever, daemon=True).start()
js_errors = []
shot = os.path.join(tmp_dir, "reqs.png")
try:
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 900})
        page.on("pageerror", lambda e: js_errors.append(str(e)))
        page.goto(f"http://127.0.0.1:{port}/login")
        page.fill("input[name=email]", EMAIL)
        page.fill("input[name=password]", PW)
        page.click("button[type=submit]")
        page.wait_for_url("**/dashboard**")
        check("app.js?v=5.81.0" in page.content(), "cache-buster not bumped to 5.81.0")
        page.wait_for_function("typeof state !== 'undefined' && state.consultants && state.consultants.length", timeout=15000)
        page.evaluate(f"openTracker({cid})")
        page.click("#trk-open-reqs")
        page.wait_for_selector("#rq-body tr[data-id]", timeout=10000)
        txt = page.inner_text("#rq-body")
        check("Virtusa is Hiring" in txt and "Draft created" in txt and "78%" in txt and "Lunch" not in txt, f"drafts shown (default filter): {txt[:300]}")
        check(page.locator("#rq-body a[href^='/api/optimized-resumes/']").count() >= 1, "tailored resume link")
        page.select_option("#rq-status", "")
        page.wait_for_function("document.getElementById('rq-body').innerText.includes('Lunch')", timeout=5000)
        check("Low match" in page.inner_text("#rq-body") and "Not a requirement" in page.inner_text("#rq-body"), "all results")
        page.screenshot(path=shot)
        browser.close()
finally:
    server.shutdown()
check(not js_errors, f"JavaScript errors: {js_errors}")

if failures:
    print(f"FAIL: {len(failures)} check(s) failed:")
    for f in failures:
        print("  -", f)
    print("screenshot:", shot)
    sys.exit(1)
print(f"PASS: requirements from Gmail - Primary only, read-only PEEK, skips (notifications / personal / replies) without wasted AI, 60% gate, tailored version attached, threaded draft to Reply-To, tracker, no duplicates, retry, background scan, privacy, window. Screenshot: {shot}")
